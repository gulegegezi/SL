"""把人生条漫的剧情导出成 Word（同时出一份 Markdown 便于对照）。

数据来源：
  data/story_life_comic.json  事件 / 四档镜头 / 选项 / 回响
  data/story_life_dialogue.json  选项前的旁白与选项后的回应
"""
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
EVENTS = ROOT / "data" / "story_life_comic.json"
DIALOGUE = ROOT / "data" / "story_life_dialogue.json"
OUT_DOCX = ROOT / "docs" / "人生条漫-剧情全量.docx"
OUT_MD = ROOT / "docs" / "人生条漫-剧情全量.md"

CN_FONT = "Microsoft YaHei"
INK = RGBColor(0x1B, 0x1B, 0x1B)
ACCENT = RGBColor(0x8C, 0x1D, 0x18)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)

BANDS = [(0, 9, "0–9 岁 · 从前有块毯子"),
         (10, 19, "10–19 岁 · 教室、停电和理想"),
         (20, 29, "20–29 岁 · 鸡蛋、绿植和电费"),
         (30, 200, "30 岁以后 · 长椅与相册")]


def style_run(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = CN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def para(doc, text, size=10.5, bold=False, color=INK, space_after=4, italic=False):
    p = doc.add_paragraph()
    style_run(p.add_run(text), size, bold, color, italic)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    return p


def collect():
    data = json.loads(EVENTS.read_text(encoding="utf-8"))
    book = json.loads(DIALOGUE.read_text(encoding="utf-8"))["events"]
    events = sorted(data["events"], key=lambda e: (e["age_range"][0], e["id"]))
    return events, book, data


def stats(events):
    chars = 0
    for e in events:
        chars += len(e.get("facts", ""))
        for lens in e.get("lenses", []):
            chars += sum(len(str(lens.get(k, ""))) for k in ("narrator", "inner", "dialogue", "result_tail"))
        for c in e.get("choices", []):
            chars += len(c.get("label", "")) + sum(len(str(x)) for x in c.get("lines", []))
    return chars


def build_md(events, book):
    lines = ["# 人生条漫 · 剧情全量", "", f"导出：{date.today().isoformat()} · {len(events)} 个事件", ""]
    for lo, hi, label in BANDS:
        group = [e for e in events if lo <= e["age_range"][0] <= hi]
        if not group:
            continue
        lines.append(f"## {label}")
        lines.append("")
        for e in group:
            lines.append(f"### {e['id']} · {e['title']}（{e['age_range'][0]}–{e['age_range'][1]} 岁 · {e.get('scene','')}）")
            lines.append("")
            lines.append(f"**开场旁白**：{e.get('facts','')}")
            if e.get("npc_dialogue"):
                lines.append(f"**人物对白**：{e['npc_dialogue']}")
            lines.append("")
            lines.append("| 档位 | 旁白 | 内心 | 台词 | 收尾 |")
            lines.append("| --- | --- | --- | --- | --- |")
            for i, lens in enumerate(e.get("lenses", [])):
                lines.append(f"| T{i} | {lens.get('narrator','')} | {lens.get('inner','')} | {lens.get('dialogue','')} | {lens.get('result_tail','')} |")
            lines.append("")
            for i, c in enumerate(e.get("choices", [])):
                resp = book.get(e["id"], {}).get("responses", [])
                r = resp[i] if i < len(resp) else ""
                roast = c.get("roast", "")
                lines.append(f"- **选项 {i+1}：{c.get('label','')}** → {r}" + (f"（旁白吐槽：{roast}）" if roast else ""))
            if e.get("echo"):
                lines.append(f"- 回响：{e['echo'].get('line','')}（来自 {e['echo'].get('from','')}）")
            lines.append("")
    return "\n".join(lines)


def build_docx(events, book, data):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(title.add_run("人生条漫 · 剧情全量"), 24, True, INK)
    sub = doc.add_paragraph()
    style_run(sub.add_run("所有事件 · 四档镜头 · 选项回应 · 旁白吐槽 · 回响"), 11, False, ACCENT)
    meta = doc.add_paragraph()
    style_run(meta.add_run(f"导出日期 {date.today().isoformat()}   ·   {len(events)} 个事件   ·   约 {stats(events)} 字"),
              9, False, MUTED)

    doc.add_paragraph()
    h = doc.add_paragraph()
    style_run(h.add_run("说明"), 13, True, ACCENT)
    para(doc, "本文件由 data/story_life_comic.json 与 data/story_life_dialogue.json 直接生成，"
              "与游戏内文本一致。前 30 岁已改成「旁白有立场」的抽象叙述：开场旁白带系统提示，"
              "选项前的旁白会自嘲，离谱选项在结果页被吐槽。")
    para(doc, "结构：每页先给开场旁白与人物对白，再用表格列四档镜头（旁白 / 内心 / 台词 / 收尾），"
              "最后是三个选项及对应回应。")

    # 目录
    doc.add_paragraph()
    h = doc.add_paragraph()
    style_run(h.add_run("分卷"), 13, True, ACCENT)
    for lo, hi, label in BANDS:
        group = [e for e in events if lo <= e["age_range"][0] <= hi]
        if group:
            para(doc, f"{label}（{len(group)} 个事件：" + "、".join(e["id"] for e in group) + "）", 10, False, MUTED)

    for lo, hi, label in BANDS:
        group = [e for e in events if lo <= e["age_range"][0] <= hi]
        if not group:
            continue
        doc.add_page_break()
        band = doc.add_paragraph()
        style_run(band.add_run(label), 18, True, INK)

        for e in group:
            head = doc.add_paragraph()
            head.paragraph_format.space_before = Pt(14)
            style_run(head.add_run(f"{e['id']} · {e['title']}"), 14, True, INK)
            sub = doc.add_paragraph()
            style_run(sub.add_run(f"{e['age_range'][0]}–{e['age_range'][1]} 岁   ·   {e.get('scene', '')}   ·   {e.get('stage', '')}"),
                      9, False, MUTED)

            p = doc.add_paragraph()
            style_run(p.add_run("开场旁白  "), 9, True, ACCENT)
            style_run(p.add_run(e.get("facts", "")), 10.5)
            p.paragraph_format.space_after = Pt(3)

            if e.get("npc_dialogue"):
                p = doc.add_paragraph()
                style_run(p.add_run("人物对白  "), 9, True, ACCENT)
                style_run(p.add_run(str(e["npc_dialogue"])), 10.5)
                p.paragraph_format.space_after = Pt(6)

            rows = e.get("lenses", [])
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            widths = [Cm(1.2), Cm(5.4), Cm(4.6), Cm(2.6), Cm(3.4)]
            hdr = ["档", "旁白", "内心", "台词", "收尾"]
            for i, text in enumerate(hdr):
                cell = table.rows[0].cells[i]
                cell.text = ""
                style_run(cell.paragraphs[0].add_run(text), 9, True, ACCENT)
            for i, lens in enumerate(rows):
                cells = table.add_row().cells
                vals = [f"T{i}", lens.get("narrator", ""), lens.get("inner", ""),
                        lens.get("dialogue", ""), lens.get("result_tail", "")]
                for j, text in enumerate(vals):
                    cells[j].text = ""
                    style_run(cells[j].paragraphs[0].add_run(str(text)), 9)
            for row in table.rows:
                for i, cell in enumerate(row.cells):
                    cell.width = widths[i]

            doc.add_paragraph()
            p = doc.add_paragraph()
            style_run(p.add_run("选项与回应"), 10, True, ACCENT)
            resp = book.get(e["id"], {}).get("responses", [])
            for i, c in enumerate(e.get("choices", [])):
                r = resp[i] if i < len(resp) else ""
                line = doc.add_paragraph(style="List Bullet")
                style_run(line.add_run(f"{c.get('label', '')}："), 10, True, INK)
                if r:
                    style_run(line.add_run(r), 10, False, INK)
                if c.get("roast"):
                    style_run(line.add_run(f"（旁白：{c['roast']}）"), 9.5, False, ACCENT)

            if e.get("echo"):
                p = doc.add_paragraph()
                style_run(p.add_run(f"回响：{e['echo'].get('line', '')}（来自 {e['echo'].get('from', '')}）"), 9, False, MUTED)

    doc.add_page_break()
    last = doc.add_paragraph()
    style_run(last.add_run("附：导出说明"), 13, True, ACCENT)
    para(doc, "重新生成：python tools/export_story_docx.py（数据改完再跑一次即可）。")
    para(doc, "校验：python tools/check_story.py；godot --headless --path . --script res://test/story_tone_smoke.gd")
    doc.save(OUT_DOCX)


def main() -> None:
    events, book, data = collect()
    OUT_MD.write_text(build_md(events, book), encoding="utf-8")
    build_docx(events, book, data)
    print("written", OUT_DOCX, OUT_MD)


if __name__ == "__main__":
    main()
